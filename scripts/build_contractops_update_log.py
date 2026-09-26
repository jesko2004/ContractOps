from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = "000000"
WHITE = "FFFFFF"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "666666"
BODY_FONT = "Microsoft YaHei"
LATIN_FONT = "Aptos"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_GRAY, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def set_repeat_heading(paragraph) -> None:
    set_keep_with_next(paragraph)
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = p_pr.find(qn("w:keepLines"))
    if keep_lines is None:
        keep_lines = OxmlElement("w:keepLines")
        p_pr.append(keep_lines)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str = BLACK) -> None:
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MID_GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for element in (fld_begin, instr_text, fld_separate, text, fld_end):
        run._r.append(element)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.34)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.18

    title = doc.styles["Title"]
    title.font.name = LATIN_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading_sizes = {"Heading 1": 16, "Heading 2": 13, "Heading 3": 11.5}
    for style_name, size in heading_sizes.items():
        style = doc.styles[style_name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 16)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_paragraph.add_run("ContractOps 项目更新记录  版本 0.5")
    set_run_font(run, size=8.5, color=BLACK)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    doc.core_properties.title = "ContractOps 项目更新记录"
    doc.core_properties.subject = "项目更新决策 技术选型 难点和验证记录"
    doc.core_properties.author = "ContractOps Project"
    doc.core_properties.keywords = "ContractOps 项目日志 技术选型 后端服务"


def add_paragraph(doc: Document, text: str = "", bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, bold=True)
        body = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(body)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.space_after = Pt(2.5)
        run = paragraph.add_run(item)
        set_run_font(run)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        run = paragraph.add_run(f"{index}.  {item}")
        set_run_font(run)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    for run in paragraph.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 11.5}[level], bold=True)
    set_repeat_heading(paragraph)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    font_size: float = 9.2,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[index])

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE if len(headers) <= 4 else PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if column_index == 0 and len(headers) > 2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            set_run_font(run, size=font_size)
            if widths:
                cell.width = Inches(widths[column_index])

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_cover(doc: Document) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("ContractOps 项目更新记录")
    set_run_font(run, size=25, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("持续记录技术选择 实现范围 问题难点和验证结果")
    set_run_font(run, size=12, color=BLACK)

    add_table(
        doc,
        ["项目字段", "当前内容"],
        [
            ["项目名称", "ContractOps 企业合同审批与履约风控后端服务"],
            ["文档版本", "0.5"],
            ["最后更新", "2026 年 9 月 26 日"],
            ["维护方式", "每次源代码更新后追加一条更新记录 不覆盖历史决策"],
            ["主要读者", "项目开发者 代码评审者 面试准备人员"],
        ],
        widths=[1.45, 5.25],
        font_size=9.8,
    )

    add_paragraph(
        doc,
        "本文件用于解释每一次项目更新为什么发生 采用了哪些候选方案 最终如何选择 "
        "实现过程中遇到什么难点 以及使用哪些证据确认修改有效",
    )
    add_paragraph(
        doc,
        "当前结论  项目主线是合同审批与履约风控后端服务 "
        "第一阶段优先验证多租户合同台账 审批状态机 履约调度 可靠事件和审计链路",
        bold_prefix="当前结论",
    )


def add_foundation_sections(doc: Document) -> None:
    add_heading(doc, "一 文档使用方法", 1)
    add_paragraph(
        doc,
        "每次修改业务边界 数据模型 服务接口 消息流程 安全策略或关键依赖时 都应新增一条记录 "
        "小型格式修复可以合并记录 但不能删除已经影响设计判断的历史内容",
    )
    add_numbered(
        doc,
        [
            "更新开始前写清问题和成功标准 避免先选技术再寻找用途",
            "列出至少一个备选方案 并记录放弃原因",
            "区分已经实现 计划实现和暂不实现的能力",
            "记录源文件 数据库迁移 配置和测试的实际变化",
            "写明未完成验证 环境限制和后续风险",
        ],
    )

    add_heading(doc, "二 更新记录字段", 1)
    add_table(
        doc,
        ["字段", "填写内容", "判断标准"],
        [
            ["更新编号", "使用 UPD 加三位数字", "编号连续且不可复用"],
            ["目标", "本次要解决的具体问题", "可以通过测试或交付物确认"],
            ["候选方案", "比较过的方法和架构", "至少说明一个未采用方案"],
            ["选择依据", "范围 成本 风险 性能和维护性", "说明真正影响决定的约束"],
            ["实现内容", "修改的模块 数据表 接口和配置", "区分完成和计划"],
            ["问题难点", "遇到的设计或实现问题", "说明原因和处理方式"],
            ["验证证据", "单测 集成测试 压测或人工检查", "不能把未执行的测试写成通过"],
            ["遗留事项", "尚未解决的风险和下一步", "给出明确优先级"],
        ],
        widths=[1.15, 3.45, 2.1],
    )

    add_heading(doc, "三 更新总览", 1)
    add_table(
        doc,
        ["编号", "日期", "主题", "结果", "验证状态"],
        [
            [
                "UPD 001",
                "2026 09 25",
                "从 EduMind 收缩为 ContractOps 并建立后端骨架",
                "完成项目定位 文档 服务骨架 数据模型和本地编排",
                "静态检查和领域逻辑已验证 完整运行待补",
            ],
            [
                "UPD 002",
                "2026 09 25",
                "调整为企业合同审批与履约风控后端服务",
                "重写业务边界 服务架构 八周计划和优先任务",
                "文档一致性和结构已检查 运行验证待后续实现",
            ],
            [
                "UPD 003",
                "2026 09 25",
                "完成 M0 工程基线",
                "加入 Alembic 请求上下文 错误协议 测试编排和专用 CI",
                "Ruff mypy 和 24 项单测通过 数据库往返待 CI",
            ],
            [
                "UPD 004",
                "2026 09 25",
                "修复上游 CI 脚本执行权限失败",
                "改为由 Bash 显式调用并消除对 Git 可执行位的依赖",
                "GitHub 自检已通过 完整检查暴露既有 Prettier 问题",
            ],
            [
                "UPD 005",
                "2026 09 26",
                "清理模型网关格式债并建立阶段 PR 审批流程",
                "格式化日志命中的八个文件并改用独立分支交付",
                "完整 Prettier 检查通过 PR CI 待运行",
            ],
        ],
        widths=[0.7, 0.9, 2.0, 2.1, 1.35],
        font_size=8.8,
    )

    add_heading(doc, "四 长期选择原则", 1)
    add_bullets(
        doc,
        [
            "先完成合同提交 审批 生效 履约 风险关闭的业务闭环 再增加模型能力",
            "领域规则保留在纯 Python 模块中 避免与 FastAPI 数据库或模型 SDK 绑定",
            "PDF 解析和模型输出只提供证据或草稿 不能直接改变合同业务状态",
            "异步任务默认会重复投递 因此消费者和副作用接口必须幂等",
            "多租户隔离同时依赖应用层权限和 PostgreSQL RLS 不能只依赖请求参数",
            "简历和演示只使用仓库内能够复现的评测和压测结果",
        ],
    )


def add_update_001(doc: Document) -> None:
    add_heading(doc, "五 更新记录 UPD 001", 1)
    add_heading(doc, "五一 更新目标", 2)
    add_paragraph(
        doc,
        "本次更新将原本覆盖 RAG 多 Agent 培训 模型网关和完整企业平台的方案收缩为 "
        "合同审查 审批和履约服务 并建立可以继续实现的后端工程骨架",
    )

    add_heading(doc, "五二 背景和范围收缩", 2)
    add_paragraph(
        doc,
        "原 EduMind 方案同时包含知识库 文档解析 多 Agent 课程生成 模型路由 可观测和部署 "
        "个人项目很难在这些方向上同时达到可运行 可测试和可量化的深度 因此本次选择一个办公场景 "
        "把后端服务能力放在主要位置",
    )
    add_paragraph(
        doc,
        "合同业务包含文档版本 异步解析 多角色审批 到期调度 权限隔离和审计等稳定需求 "
        "RAG 和模型结构化输出可以服务于明确流程 而不是单独构成产品",
    )

    add_heading(doc, "五三 垂直方向选择过程", 2)
    add_table(
        doc,
        ["候选方向", "后端深度", "个人可完成性", "办公适配", "主要问题"],
        [
            ["通用知识培训平台", "高", "低", "中", "范围过宽 容易只有目录和演示界面"],
            ["SRE 故障演练", "高", "中低", "中", "需要可信的故障仿真和工具环境"],
            ["安全 SOC 演练", "高", "低", "中", "数据和安全场景构造成本较高"],
            ["采购和招投标", "中高", "中高", "高", "文档差异较大 业务规则需要额外准备"],
            ["合同审查和履约", "高", "高", "高", "法律风险要求明确的人为审批边界"],
        ],
        widths=[1.55, 0.9, 1.1, 0.85, 2.45],
        font_size=8.7,
    )
    add_paragraph(
        doc,
        "最终选择合同方向的主要原因是业务对象清晰 状态变化可建模 后端难点真实 "
        "同时可以把模型能力限制在条款提取 风险提示和证据检索范围内",
    )

    add_heading(doc, "五四 架构选择过程", 2)
    add_paragraph(
        doc,
        "第一阶段采用模块化单体 API 加独立 Worker 不直接拆成多个业务微服务 "
        "合同 版本 审查 审批和义务属于同一业务边界 过早拆分会增加事务和部署成本 "
        "文档解析与模型调用耗时长且需要独立重试 因此计划作为异步 Worker 运行",
    )
    add_table(
        doc,
        ["决策项", "选择", "未选择方案", "选择原因", "状态"],
        [
            ["服务形态", "模块化单体加 Worker", "多业务微服务", "保持事务边界清晰 降低部署成本", "已决定"],
            ["API 框架", "FastAPI", "首版使用 Go", "便于接入解析 RAG 和模型生态", "已建立骨架"],
            ["主数据库", "PostgreSQL 加 pgvector", "PostgreSQL 加 OpenSearch", "首版同时承担事务 全文和向量检索", "已建表"],
            ["对象存储", "MinIO", "数据库保存文件", "合同原文件与结构化数据分离", "已配置"],
            ["任务队列", "Redis Streams", "Kafka 或 Celery", "支持消费组和重放且本地成本较低", "计划实现"],
            ["模型路由", "Switchyard 可选内部依赖", "模型网关承担业务权限", "业务鉴权保留在 Contract API", "保留现有代码"],
            ["租户隔离", "JWT 上下文加 RLS", "仅使用 tenant_id 查询条件", "应用错误时数据库仍提供第二道隔离", "数据库已准备"],
            ["事件一致性", "Transactional Outbox", "事务后直接发布消息", "避免数据库成功但事件丢失", "表结构已准备"],
        ],
        widths=[1.05, 1.35, 1.35, 2.35, 0.75],
        font_size=8.3,
    )

    add_heading(doc, "五五 本次源文件变化", 2)
    add_table(
        doc,
        ["区域", "主要文件", "变化"],
        [
            ["项目入口", "README CONTRACTOPS 和根目录 README", "主线改为 ContractOps 并保留 OpenMAIC 为演示端"],
            ["范围文档", "CONTRACTOPS SCOPE", "定义 MVP 已实现 未实现和明确不做的内容"],
            ["计划书", "CONTRACTOPS IMPLEMENTATION PLAN", "加入架构 领域模型 API 事件 八周计划和测试矩阵"],
            ["API 骨架", "services contract api src", "加入配置 应用工厂 健康检查 系统信息和请求 ID"],
            ["领域规则", "domain review py", "实现审查状态枚举 合法迁移和异常"],
            ["数据库", "migrations 0001 initial sql", "建立合同 版本 任务 Chunk 审查 审批 义务 Outbox 和审计表"],
            ["本地环境", "deploy contractops docker compose yml", "配置 API PostgreSQL Redis MinIO 和独立应用角色"],
            ["测试", "tests test review state machine 和 test health", "覆盖状态迁移 健康检查和请求 ID"],
        ],
        widths=[1.1, 2.2, 3.6],
        font_size=8.7,
    )

    add_heading(doc, "五六 问题难点和处理", 2)

    add_heading(doc, "范围控制和上游代码保留", 3)
    add_paragraph(
        doc,
        "仓库已经包含完整 OpenMAIC 和 Switchyard 代码 如果直接删除会丢失可用的演示能力和许可证信息 "
        "如果继续沿用原文档又会让读者误以为平台能力已经完成 处理方式是新增 ContractOps 入口 "
        "将旧 EduMind 文档改为迁移指引 保留上游源码但降低其业务地位",
    )

    add_heading(doc, "RLS 对表所有者默认不生效", 3)
    add_paragraph(
        doc,
        "最初 Compose 使用同一个 PostgreSQL 用户初始化数据库并连接 API 该用户作为表所有者会绕过普通 RLS "
        "这会使多租户策略看起来存在但实际无法提供隔离 最终改为 contractops admin 执行迁移 "
        "contractops app 运行服务 并对业务表启用 FORCE ROW LEVEL SECURITY",
    )

    add_heading(doc, "允许空值的唯一约束", 3)
    add_paragraph(
        doc,
        "合同版本上传完成前没有 content hash 如果使用 NULLS NOT DISTINCT "
        "同一合同只能创建一个尚未完成上传的版本 本次改为 content hash 非空时才生效的部分唯一索引 "
        "既允许并发上传草稿 又能在哈希确定后阻止重复版本",
    )

    add_heading(doc, "当前版本指针形成循环外键", 3)
    add_paragraph(
        doc,
        "contracts 需要指向当前版本 contract versions 又必须属于某个合同 "
        "迁移先创建两个表 再补充 current version 复合外键 并将该约束设为可延迟 "
        "业务写入时仍采用先创建合同 再创建版本 最后切换当前版本的顺序",
    )

    add_heading(doc, "不能把未运行测试写成通过", 3)
    add_paragraph(
        doc,
        "当前环境缺少 FastAPI pytest 和 Docker 因此无法执行完整 API 单测 数据库迁移和 Compose 启动 "
        "本次只记录已完成的 Python AST 解析 TOML 解析 Markdown 结构检查和纯领域状态机验证 "
        "完整运行验证保留为下一阶段工作",
    )

    add_heading(doc, "大批量补丁不利于定位错误", 3)
    add_paragraph(
        doc,
        "一次性替换旧文档并新增大量文件时 补丁工具拒绝同一路径的多次操作 "
        "随后将修改拆为项目文档 服务源码 数据库和部署配置四个批次 每个批次完成后检查目标文件 "
        "以后更新也沿用小批次写入和逐步验证",
    )

    add_heading(doc, "五七 已执行验证", 2)
    add_table(
        doc,
        ["验证项", "结果", "说明"],
        [
            ["Python AST 解析", "通过", "contract api 下所有 Python 文件语法可解析"],
            ["pyproject TOML 解析", "通过", "项目元数据和工具配置可读取"],
            ["审查状态机直接验证", "通过", "合法迁移成功 非法迁移抛出异常"],
            ["Markdown 围栏检查", "通过", "核心文档代码围栏成对"],
            ["生成缓存清理", "通过", "未遗留 pycache 目录"],
            ["FastAPI 与 pytest 测试", "未执行", "当前 Python 环境未安装依赖"],
            ["数据库迁移", "未执行", "当前环境没有 Docker 和 PostgreSQL"],
            ["Compose 启动", "未执行", "当前环境没有 Docker"],
        ],
        widths=[1.7, 1.0, 4.2],
        font_size=9,
    )

    add_heading(doc, "五八 遗留风险", 2)
    add_bullets(
        doc,
        [
            "迁移 SQL 仍需在真实 pgvector PostgreSQL 上执行并验证 RLS 策略",
            "API 尚未实现认证 数据库事务和 SET LOCAL app tenant id",
            "MinIO 镜像当前用于本地开发 正式环境需要固定版本并配置密钥管理",
            "中文合同的全文检索效果不能只依赖 PostgreSQL simple 配置 需要评测分词方案",
            "模型审查属于辅助意见 需要低置信度和高风险条款的强制人工复核",
            "履约提醒尚未实现租约 去重 重试和取消传播",
        ],
    )

    add_heading(doc, "五九 下一步", 2)
    add_numbered(
        doc,
        [
            "补充 Alembic 或等价迁移管理 并在空数据库执行初始迁移",
            "实现 JWT 租户上下文 RBAC 和数据库事务内的租户设置",
            "实现 Contract 和 ContractVersion Repository 及创建接口",
            "实现 MinIO 预签名上传和完成上传幂等接口",
            "实现 Outbox Publisher 和 Redis Streams Worker 骨架",
        ],
    )


def add_update_002(doc: Document) -> None:
    add_heading(doc, "六 更新记录 UPD 002", 1)

    add_heading(doc, "六一 更新目标", 2)
    add_paragraph(
        doc,
        "本次更新解决 ContractOps 与 PDF Inspector 可能被理解为同类项目的问题 "
        "将产品主线明确为企业合同审批与履约风控后端服务 并重新安排实现顺序",
    )

    add_heading(doc, "六二 候选定位比较", 2)
    add_table(
        doc,
        ["候选定位", "主要能力", "后端深度", "与 PDF Inspector 重合", "决定"],
        [
            ["合同文档审查工具", "解析 条款提取 风险提示", "中", "高", "不采用"],
            ["合同全生命周期平台", "起草 审批 签署 履约 结算", "高", "低", "范围过大"],
            ["合同审批与履约风控后端", "审批 待办 义务 风险 提醒 审计", "高", "低", "采用"],
        ],
        widths=[1.55, 2.1, 0.85, 1.25, 0.85],
        font_size=8.7,
    )
    add_paragraph(
        doc,
        "最终选择第三个方案 PDF 和模型能力保留为可关闭的辅助模块 "
        "即使不启用文档解析 用户仍可以通过结构化字段完成合同审批和履约管理",
    )

    add_heading(doc, "六三 新的业务主链路", 2)
    add_numbered(
        doc,
        [
            "合同经办人创建台账并提交不可变版本",
            "系统按照合同类型 金额 部门和风险标签匹配审批策略",
            "法务 财务和业务角色处理个人待办并形成可追溯决定",
            "合同生效后登记付款 交付 验收 续约和终止通知等义务",
            "调度器生成到期或逾期风险事件并发送去重提醒",
            "风险负责人完成确认 延期 升级或关闭并保留审计记录",
        ],
    )

    add_heading(doc, "六四 架构重点变化", 2)
    add_table(
        doc,
        ["原重点", "调整后重点", "原因", "实现顺序"],
        [
            ["AI 审查状态机", "合同生命周期加审批实例状态机", "业务状态不应由模型任务驱动", "优先"],
            ["RAG 与风险 Finding", "审批策略 待办和策略快照", "突出企业流程和并发控制", "优先"],
            ["义务提取", "义务台账 数据库租约和风险事件", "没有模型也能完成履约闭环", "优先"],
            ["文档解析 Worker", "可靠事件 通知和调度 Worker", "先验证幂等 恢复和消息一致性", "优先"],
            ["PDF 页码引用", "可选文档证据适配器", "降低与 PDF Inspector 的重合", "第七周"],
            ["模型路由", "可关闭的智能适配器", "模型不能阻塞主业务闭环", "后置"],
        ],
        widths=[1.35, 2.1, 2.25, 1.0],
        font_size=8.6,
    )

    add_heading(doc, "六五 本次源文件变化", 2)
    add_table(
        doc,
        ["文件", "修改内容", "完成状态"],
        [
            ["README CONTRACTOPS", "更新项目名称 当前能力和未实现范围", "已完成"],
            ["CONTRACTOPS SCOPE", "增加 PDF Inspector 边界 新闭环 核心能力和成功标准", "已完成"],
            ["CONTRACTOPS IMPLEMENTATION PLAN", "重写架构 领域模型 API 事件 八周计划和测试矩阵", "已完成"],
            ["项目更新记录", "新增 UPD 002 决策过程 风险和下一步", "已完成"],
            ["运行时代码和数据库迁移", "本次不修改 等范围稳定后按新计划实施", "未开始"],
        ],
        widths=[2.0, 3.8, 1.1],
        font_size=8.8,
    )

    add_heading(doc, "六六 关键难点", 2)

    add_heading(doc, "避免只修改名称而保留旧实现顺序", 3)
    add_paragraph(
        doc,
        "如果只把项目标题改成审批与履约风控 但仍先实现 PDF 解析 RAG 和模型审查 "
        "最终演示仍会像文档检查工具 因此本次同时调整领域模型 API 事件和八周里程碑 "
        "把审批策略 可靠事件和履约调度放到文档智能之前",
    )

    add_heading(doc, "合同状态和技术任务状态不能混用", 3)
    add_paragraph(
        doc,
        "旧方案把 INGESTING 和 AI REVIEW 放进合同审查主状态机 "
        "新的设计把合同生命周期 审批步骤状态 文件解析状态拆开保存 "
        "解析失败只影响辅助证据 不直接改变已经存在的合同业务状态",
    )

    add_heading(doc, "审批策略需要版本快照", 3)
    add_paragraph(
        doc,
        "企业管理员可能在审批进行中修改策略 如果审批实例每次读取最新规则 "
        "同一合同前后步骤会使用不同条件 新方案要求发布策略版本并在创建审批实例时保存快照 "
        "已经开始的流程不会被后续配置变更影响",
    )

    add_heading(doc, "履约风险不能等同于自动违约判断", 3)
    add_paragraph(
        doc,
        "调度器可以根据日期产生到期或逾期风险 但不能自动认定合同违约 "
        "系统只创建风险事件并要求负责人确认 延期 升级或关闭 "
        "最终处置保持明确的人为责任边界",
    )

    add_heading(doc, "六七 验证和限制", 2)
    add_table(
        doc,
        ["验证项", "结果", "说明"],
        [
            ["范围一致性检查", "通过", "README 范围文档和实施计划使用相同定位"],
            ["PDF Inspector 边界检查", "通过", "通用 PDF 检查已列入明确不做"],
            ["里程碑依赖检查", "通过", "审批和履约先于文档解析和模型能力"],
            ["代码运行验证", "未执行", "本次只调整方案 尚未修改运行时代码"],
            ["Word 结构检查", "通过", "标题 表格 更新编号和不可跨页行已检查"],
            ["Word 最终视觉渲染", "通过", "使用临时解包渲染器逐页检查 不安装系统组件"],
        ],
        widths=[1.7, 1.0, 4.2],
        font_size=9,
    )

    add_heading(doc, "六八 下一步", 2)
    add_numbered(
        doc,
        [
            "将初始迁移纳入 Alembic 或等价管理并补空库迁移测试",
            "实现 JWT 请求上下文 RBAC 部门数据范围和事务级 RLS",
            "实现 Contract 与 ContractVersion Repository 和创建查询接口",
            "新增 ApprovalPolicy ApprovalInstance 和 ApprovalStep 数据模型",
            "实现审批策略匹配 策略快照 提交审批和个人待办接口",
            "实现审批决定的幂等 乐观锁 审计事件和 Outbox 原子写入",
        ],
    )


def add_update_003(doc: Document) -> None:
    add_heading(doc, "七 更新记录 UPD 003", 1)

    add_heading(doc, "七一 更新目标", 2)
    add_paragraph(
        doc,
        "本次更新执行实施计划的 M0 工程基线 建立后续业务开发必须复用的应用入口 "
        "错误协议 请求上下文 数据库迁移 测试环境和持续集成入口",
    )

    add_heading(doc, "七二 方案选择过程", 2)
    add_table(
        doc,
        ["决策项", "候选方案", "最终选择", "选择依据"],
        [
            ["迁移管理", "继续使用容器初始化 SQL 或改用 ORM 自动建表", "Alembic 调用已评审 SQL", "保留 RLS 部分索引和 pgvector 等 PostgreSQL 能力 同时获得版本和回滚"],
            ["请求上下文", "BaseHTTPMiddleware 或纯 ASGI", "纯 ASGI 加 ContextVar", "避免中间件任务边界造成上下文传播差异 并为 M1 租户身份预留入口"],
            ["错误协议", "各接口自行返回或全局处理", "全局稳定错误信封", "统一错误码 消息 请求 ID 和可选详情 方便客户端和审计关联"],
            ["数据库启动", "PostgreSQL 初始化目录直接执行建表 SQL", "一次性 migrate 服务", "迁移版本成为唯一事实来源 API 只在迁移成功后启动"],
            ["持续集成", "塞入上游大型 CI 或单独建立门禁", "路径过滤的 ContractOps API CI", "缩短反馈时间并避免 ContractOps 检查依赖整个 OpenMAIC 构建"],
        ],
        widths=[1.15, 1.75, 1.65, 2.45],
        font_size=8.3,
    )

    add_heading(doc, "七三 已完成实现", 2)
    add_table(
        doc,
        ["区域", "实现内容", "主要文件"],
        [
            ["应用基线", "可注入 Settings 的应用工厂 纯 ASGI 请求上下文和请求 ID", "main py context py request context py"],
            ["错误处理", "领域异常 参数错误 HTTP 错误和未知错误使用同一响应结构", "errors py 和 test errors py"],
            ["领域规则", "实现计划书中的合同生命周期和非法迁移拒绝", "domain contract py 和状态机测试"],
            ["迁移管理", "Alembic 基线 升级 SQL 回滚 SQL和重复升级回滚验证脚本", "alembic ini migrations 和 verify migrations py"],
            ["本地编排", "开发 Compose 使用 migrate 服务 新增隔离测试 Compose", "deploy contractops 下两个 Compose 文件"],
            ["持续集成", "加入 Ruff mypy pytest 空库迁移和 Compose 配置检查", "contractops api yml"],
            ["稳定入口", "一个命令运行静态检查 类型检查 单测和可选迁移验证", "scripts check py"],
            ["项目说明", "同步根说明中的 M0 已完成能力 启动顺序和统一检查命令", "README CONTRACTOPS"],
        ],
        widths=[1.2, 3.45, 2.35],
        font_size=8.5,
    )

    add_heading(doc, "七四 关键难点", 2)

    add_heading(doc, "迁移需要兼容 PostgreSQL 专有结构", 3)
    add_paragraph(
        doc,
        "现有建表语句包含 pgvector HNSW 索引 生成列 延迟外键和动态 RLS 策略 "
        "如果立即改写为 ORM 元数据 容易遗漏约束或改变执行语义 本次让 Alembic 管理经过评审的 SQL 文件 "
        "后续迁移仍使用明确的升级和回滚脚本",
    )

    add_heading(doc, "运行角色不能承担迁移权限", 3)
    add_paragraph(
        doc,
        "扩展创建 表结构和 RLS 策略需要高于应用运行角色的权限 因此配置单独的 "
        "CONTRACTOPS MIGRATION DATABASE URL Compose 中由短生命周期 migrate 服务使用 "
        "常驻 API 继续使用 contractops app",
    )

    add_heading(doc, "请求 ID 必须覆盖错误响应", 3)
    add_paragraph(
        doc,
        "请求 ID 如果只在正常响应后添加 异常路径会缺少关联标识 本次中间件在 ASGI response start 阶段统一写入响应头 "
        "错误处理器从请求状态读取同一个值 因此领域错误 404 和参数错误都能返回相同请求 ID",
    )

    add_heading(doc, "当前机器没有 Docker 和 PostgreSQL", 3)
    add_paragraph(
        doc,
        "本机无法执行真实 pgvector 数据库的升级回滚循环 因此本次没有把迁移往返写成通过 "
        "已新增隔离测试 Compose 和 GitHub Actions 数据库服务 CI 会在空库执行两次升级 回滚到 base 再升级到 head",
    )

    add_heading(doc, "七五 验证结果", 2)
    add_table(
        doc,
        ["验证项", "结果", "证据或限制"],
        [
            ["Ruff", "通过", "src tests scripts migrations 全部通过"],
            ["mypy strict", "通过", "15 个 Python 源文件无类型错误"],
            ["pytest", "通过", "24 项测试通过 覆盖 API 错误 请求上下文和两个状态机"],
            ["Alembic 版本发现", "通过", "history 正确识别 20260925 0001 head"],
            ["YAML 解析", "通过", "开发 Compose 测试 Compose 和专用 CI 均可解析"],
            ["Python compileall", "通过", "src tests scripts migrations 可以编译"],
            ["真实数据库迁移往返", "未执行", "本机无 Docker 和 PostgreSQL 由新增 CI 执行"],
            ["测试 Compose 启动", "未执行", "本机无 Docker"],
        ],
        widths=[1.75, 1.0, 4.2],
        font_size=8.8,
    )

    add_heading(doc, "七六 下一步", 2)
    add_numbered(
        doc,
        [
            "在 GitHub Actions 或具备 Docker 的环境确认空库迁移和回滚测试通过",
            "进入 M1 实现 JWT 租户上下文 RBAC 部门数据范围和事务级 RLS 设置",
            "实现 Contract 与 ContractVersion Repository 以及创建 查询和新增版本接口",
            "补充两个租户互相读取和重复请求不产生重复版本的集成测试",
        ],
    )


def add_update_004(doc: Document) -> None:
    add_heading(doc, "八 更新记录 UPD 004", 1)

    add_heading(doc, "八一 更新目标", 2)
    add_paragraph(
        doc,
        "修复 GitHub Actions 中 Lint Typecheck 与 Unit Tests 任务在 Parallel runner self test 步骤立即失败的问题 "
        "使上游前端检查能够继续执行并避免 Windows 上传方式再次触发同类错误",
    )

    add_heading(doc, "八二 问题定位", 2)
    add_paragraph(
        doc,
        "公开 GitHub Actions API 显示运行 36132265036 的任务 108063921078 在第三步返回退出码 126 "
        "后续依赖安装 静态检查和单元测试因此全部跳过 本地 Git 索引显示 scripts ci run parallel sh 的模式为 100644 "
        "Ubuntu runner 直接执行该文件时没有执行权限",
    )

    add_heading(doc, "八三 方案选择", 2)
    add_table(
        doc,
        ["候选方案", "优点", "限制", "决定"],
        [
            ["将脚本模式改为 100755", "保留原工作流调用形式", "Windows 或 GitHub API 上传可能再次丢失可执行位", "不单独采用"],
            ["工作流显式使用 Bash", "不依赖文件模式 对不同上传方式稳定", "要求 Ubuntu runner 提供 Bash", "采用"],
            ["在运行前执行 chmod", "可以继续直接运行脚本", "增加无必要的工作区修改并掩盖调用约束", "不采用"],
        ],
        widths=[1.75, 2.0, 2.25, 1.0],
        font_size=8.7,
    )

    add_heading(doc, "八四 实现内容", 2)
    add_paragraph(
        doc,
        "修改 github workflows ci yml 中三处调用 在 Parallel runner self test 的成功和预期失败分支以及并行质量检查入口前增加 bash "
        "脚本内容和任务并发语义保持不变",
    )

    add_heading(doc, "八五 验证结果", 2)
    add_table(
        doc,
        ["验证项", "结果", "证据或限制"],
        [
            ["失败原因核对", "通过", "GitHub 检查注释为退出码 126 且失败发生在直接执行脚本的第一步"],
            ["成功路径自检", "通过", "Git Bash 执行两个成功命令并返回零"],
            ["预期失败路径", "通过", "内部命令退出 7 时并行运行器汇总后返回 1"],
            ["工作流差异检查", "通过", "三处调用均改为 bash scripts ci run parallel sh"],
            ["GitHub Actions 自检", "通过", "运行 36138239412 的 Parallel runner self test 已通过"],
            ["完整质量检查", "发现后续问题", "脚本权限问题消失 后续 Prettier 检查发现八个既有格式文件"],
        ],
        widths=[1.65, 1.0, 4.3],
        font_size=8.8,
    )

    add_heading(doc, "八六 后续动作", 2)
    add_numbered(
        doc,
        [
            "将后续格式修复放入独立阶段分支并创建 PR",
            "观察 PR CI 是否通过完整质量检查",
            "由项目负责人审批后再合并到远端 main",
        ],
    )


def add_update_005(doc: Document) -> None:
    add_heading(doc, "九 更新记录 UPD 005", 1)

    add_heading(doc, "九一 更新目标", 2)
    add_paragraph(
        doc,
        "处理脚本权限修复后继续暴露的 Prettier 失败 同时把阶段交付调整为独立分支和 PR 审批流程 "
        "避免后续阶段直接更新远端 main",
    )

    add_heading(doc, "九二 问题定位", 2)
    add_paragraph(
        doc,
        "GitHub Actions 运行 36138239412 已确认 Parallel runner self test 通过 "
        "随后 Prettier 3 8 1 报告 services model gateway 下八个既有 JSON CSS 和无扩展名 README 不符合根仓库格式规则 "
        "其余 E2E Render Service Issue Triage 和 Main image 均通过",
    )

    add_heading(doc, "九三 方案选择", 2)
    add_table(
        doc,
        ["候选方案", "优点", "限制", "决定"],
        [
            ["只格式化日志命中的八个文件", "改动可审查 不改变检查标准", "会产生较大的纯格式差异", "采用"],
            ["把模型网关目录加入 Prettier ignore", "改动最少", "隐藏持续存在的格式债并降低门禁覆盖", "不采用"],
            ["关闭或放宽 Prettier 门禁", "CI 可以快速变绿", "破坏仓库统一质量标准", "不采用"],
        ],
        widths=[1.75, 2.0, 2.25, 1.0],
        font_size=8.7,
    )

    add_heading(doc, "九四 实现内容", 2)
    add_bullets(
        doc,
        [
            "新建 fix ci prettier gate 阶段分支",
            "使用仓库锁定的 Prettier 3 8 1 仅格式化 CI 日志列出的八个文件",
            "保留现有 Prettier ignore 和检查命令 不降低质量门禁",
            "阶段完成后创建 PR 等待项目负责人审批 再合并到远端 main",
        ],
    )

    add_heading(doc, "九五 验证结果", 2)
    add_table(
        doc,
        ["验证项", "结果", "证据或限制"],
        [
            ["原权限故障", "通过", "GitHub Runner 上 Parallel runner self test 已通过"],
            ["完整 Prettier 检查", "通过", "Prettier 3 8 1 检查整个仓库并返回全部匹配"],
            ["ContractOps 后端门禁", "通过", "Ruff mypy strict 和 24 项 pytest 通过"],
            ["PR CI", "待执行", "分支提交和 PR 尚待创建"],
            ["合并 main", "待审批", "只有项目负责人批准后执行"],
        ],
        widths=[1.65, 1.0, 4.3],
        font_size=8.8,
    )

    add_heading(doc, "九六 后续动作", 2)
    add_numbered(
        doc,
        [
            "提交阶段分支并通过 GitHub API 推送远端分支",
            "创建 PR 并观察全部 Actions 结果",
            "向项目负责人提交 PR 链接和验证证据 等待明确审批",
        ],
    )


def add_decisions_and_risks(doc: Document) -> None:
    add_heading(doc, "十 当前决策记录", 1)
    add_table(
        doc,
        ["编号", "决策", "原因", "重新评估条件"],
        [
            ["ADR 001", "主线改为 ContractOps", "缩小范围并提高后端实现深度", "合同场景缺少可用数据或无法形成闭环"],
            ["ADR 002", "模块化单体加 Worker", "保留领域事务并隔离长任务", "检索或调度需要独立扩容"],
            ["ADR 003", "PostgreSQL 加 pgvector", "一个系统完成事务 全文和向量存储", "规模或中文召回达到明确瓶颈"],
            ["ADR 004", "Redis Streams 作为首版队列", "本地成本低 支持消费组和消息重领", "需要长期回放或更高吞吐"],
            ["ADR 005", "应用权限加 PostgreSQL RLS", "防止查询条件遗漏导致跨租户读取", "不能删除 只允许调整实现"],
            ["ADR 006", "关键法律和审批决定由人完成", "模型输出存在错误和法律责任边界", "不能删除"],
            ["ADR 007", "主线改为审批与履约风控", "避免退化为 PDF 检查工具并突出企业后端能力", "审批或履约场景无法形成可运行闭环"],
            ["ADR 008", "审批实例保存策略快照", "避免流程中途因策略变更产生不一致", "只允许优化快照存储方式"],
            ["ADR 009", "Alembic 管理数据库版本", "让升级 回滚和部署顺序可以复现", "不能退回容器初始化目录直接建业务表"],
            ["ADR 010", "统一错误信封和请求上下文", "客户端和审计可以依靠稳定错误码和请求 ID", "只允许兼容性扩展"],
            ["ADR 011", "CI Shell 脚本由 Bash 显式调用", "避免 Windows 和 API 上传丢失可执行位", "只有上传链路稳定保留 100755 时才评估简化"],
            ["ADR 012", "阶段成果使用独立分支和 PR 审批", "让 main 只接收经过检查和人工批准的变更", "只有项目负责人明确调整交付流程时变更"],
        ],
        widths=[0.8, 1.6, 2.6, 1.9],
        font_size=8.7,
    )

    add_heading(doc, "十一 风险登记", 1)
    add_table(
        doc,
        ["风险", "影响", "当前控制", "后续验证"],
        [
            ["审批策略配置错误", "错误人员审批或必要步骤被跳过", "受控条件 策略发布和实例快照", "策略模拟测试和高金额样本回放"],
            ["模型生成无依据风险", "错误提示或遗漏关键条款", "模型只生成草稿并保存版本 页码和原文", "建立引用准确率和风险召回评测集"],
            ["跨租户数据泄露", "合同和商业信息外泄", "独立应用角色 FORCE RLS 对象键带租户", "集成测试和恶意越权测试"],
            ["恶意合同内容", "Prompt 注入或解析器攻击", "模型不执行合同内指令 工具默认关闭", "恶意样本和文件沙箱测试"],
            ["重复任务和提醒", "重复审查 重复通知", "幂等表 Outbox 唯一键设计", "Worker 崩溃和重复投递测试"],
            ["解析定位错误", "引用页码不准确", "保存版本 Chunk 页码和 parser version", "人工标注样本比对"],
            ["合同修订覆盖历史", "审计链丢失", "版本不可变 驳回后创建新版本", "版本 Diff 和回滚测试"],
            ["迁移权限泄露给运行服务", "应用漏洞可能修改结构或绕过隔离", "迁移 URL 与运行 URL 分离 migrate 服务短期运行", "CI 检查运行角色权限和生产密钥配置"],
            ["跨平台文件模式丢失", "Linux CI 脚本在任务开始时退出 126", "工作流显式使用 Bash 调用脚本", "每次迁移上传方式后观察 CI 自检"],
            ["上游格式债阻断阶段交付", "业务代码通过但根仓库质量门禁失败", "按日志修复具体文件 不扩大忽略范围", "PR 中运行完整 Prettier 和现有 CI"],
        ],
        widths=[1.45, 1.55, 2.45, 1.45],
        font_size=8.7,
    )


def add_update_template(doc: Document) -> None:
    add_heading(doc, "十二 后续更新记录模板", 1)
    add_paragraph(
        doc,
        "复制本节并替换方括号内容 新记录必须追加在历史记录之后 不修改已经完成的决策说明",
    )

    add_heading(doc, "更新记录 UPD 编号", 2)
    add_table(
        doc,
        ["字段", "填写内容"],
        [
            ["日期", "填写 YYYY MM DD"],
            ["更新目标", "说明本次要解决的具体问题"],
            ["关联里程碑", "填写 M0 到 M7 或新增里程碑"],
            ["更新状态", "计划中 进行中 已完成 已回滚"],
        ],
        widths=[1.45, 5.25],
        font_size=9.5,
    )

    template_sections = [
        ("背景和问题", "说明触发本次更新的现象 约束和成功标准"),
        ("候选方案", "列出考虑过的方法及各自优缺点"),
        ("选择依据", "说明真正影响决定的范围 成本 风险 性能和维护性因素"),
        ("最终决定", "写明采用的方法和不采用其他方案的原因"),
        ("源文件变化", "列出新增 修改 删除的文件 数据库迁移 接口和配置"),
        ("问题难点", "记录错误根因 处理过程和仍不确定的部分"),
        ("验证结果", "记录实际执行的测试 命令 数据和结果"),
        ("遗留风险", "列出未验证内容 技术债和影响"),
        ("下一步", "列出下一批可以直接执行的任务"),
    ]
    for title, prompt in template_sections:
        add_heading(doc, title, 3)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)
        run = paragraph.add_run(f"填写提示  {prompt}")
        set_run_font(run, color=MID_GRAY)

    add_heading(doc, "更新完成检查", 2)
    add_bullets(
        doc,
        [
            "源代码和文档中的完成状态一致",
            "数据迁移和配置变更可以回退",
            "测试结果写明实际执行环境",
            "没有把计划能力描述为已完成",
            "新增决策和新增风险已经分别同步到当前决策记录与风险登记",
        ],
    )


def build_document(output_path: Path) -> None:
    document = Document()
    configure_document(document)
    add_cover(document)
    add_foundation_sections(document)
    add_update_001(document)
    add_update_002(document)
    add_update_003(document)
    add_update_004(document)
    add_update_005(document)
    add_decisions_and_risks(document)
    add_update_template(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("architecture/ContractOps-项目更新记录.docx"),
    )
    args = parser.parse_args()
    build_document(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()

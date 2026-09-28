from __future__ import annotations

from pathlib import Path

from build_contractops_update_log import (
    PALE_BLUE,
    add_bullets,
    add_heading,
    add_paragraph,
    add_table,
    prevent_row_split,
    set_cell_margins,
    set_cell_shading,
    set_run_font,
)
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOCUMENT_PATH = Path("architecture/ContractOps-项目更新记录.docx")


def _set_cell(cell, value: str, *, shaded: bool = False, centered: bool = False) -> None:
    cell.text = ""
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if shaded:
        set_cell_shading(cell, PALE_BLUE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(value)
    set_run_font(run, size=8.7)


def _append_table_row(table, values: list[str]) -> None:
    row = table.add_row()
    prevent_row_split(row)
    shaded = (len(table.rows) - 2) % 2 == 1
    for index, value in enumerate(values):
        _set_cell(row.cells[index], value, shaded=shaded, centered=index == 0)


def _table_with_headers(document: Document, headers: tuple[str, ...]):
    for table in document.tables:
        if tuple(cell.text.strip() for cell in table.rows[0].cells) == headers:
            return table
    raise RuntimeError(f"table not found: {headers}")


def _set_metadata(document: Document) -> None:
    for section in document.sections:
        paragraph = section.header.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("ContractOps 项目更新记录  版本 1.3")
        set_run_font(run, size=8.5)

    cover = _table_with_headers(document, ("项目字段", "当前内容"))
    values = {"文档版本": "1.3", "最后更新": "2026 年 9 月 28 日"}
    for row in cover.rows[1:]:
        key = row.cells[0].text.strip()
        if key in values:
            _set_cell(row.cells[1], values[key])

    document.paragraphs[4].text = (
        "当前结论  项目主线是合同审批与履约风控后端服务 M0 至 M7 已形成多租户台账 "
        "审批 履约 可靠事件 审计 文档证据 故障恢复和可复现性能验证的后端闭环"
    )
    for run in document.paragraphs[4].runs:
        set_run_font(run)


def _ensure_overview(document: Document) -> None:
    overview = _table_with_headers(document, ("编号", "日期", "主题", "结果", "验证状态"))
    if any(row.cells[0].text.strip() == "UPD 013" for row in overview.rows[1:]):
        return
    _append_table_row(
        overview,
        [
            "UPD 013",
            "2026 09 28",
            "实现 M7 故障验证与压测",
            "新增真实服务压测 故障恢复 安全门禁 运行手册 演示脚本和 CI 证据产物",
            "本地 62 项测试 Ruff mypy 通过 性能数字待 M7 PR CI 实跑回填",
        ],
    )


def _add_update(document: Document) -> list[object]:
    before = set(document.element.body)

    add_heading(document, "十七 更新记录 UPD013", 1)
    add_heading(document, "十七之一 更新目标", 2)
    add_paragraph(
        document,
        "执行实施计划 M7 故障验证和压测 将已有 API Repository Worker Scheduler 安全与观测能力 "
        "组合成可重复执行的生产就绪证据链 使用真实 FastAPI PostgreSQL Redis 运行环境产生性能数据 "
        "并提供故障演练 运行处置和最终项目演示方法",
    )

    add_heading(document, "十七之二 候选方案和选择过程", 2)
    add_table(
        document,
        ["决策项", "候选方法", "最终选择", "选择依据"],
        [
            [
                "性能工具",
                "自写计时脚本 k6 或 Locust",
                "Locust 无界面模式",
                "能表达带 JWT 的多接口用户流 输出 CSV HTML 并可在 Python 服务仓库直接维护",
            ],
            [
                "运行环境",
                "本机内存替身或 CI 真实依赖",
                "GitHub CI 的 Uvicorn PostgreSQL 16 Redis",
                "本机没有 Docker 不把 TestClient 速度伪装为服务 QPS CI 可以固定提交和运行 ID",
            ],
            [
                "性能门禁",
                "只展示报表或设定可执行阈值",
                "QPS P95 失败率阈值加原始产物",
                "阈值阻止明显回退 原始 CSV HTML JSON 支持复核且不把 CI 基线称为生产容量",
            ],
            [
                "故障验证",
                "手工停服务或确定性故障注入",
                "自动化注入加运行手册人工演练",
                "自动化保证回归 手册覆盖真实进程崩溃 Redis 暂停通知失败和租约恢复",
            ],
            [
                "Webhook 安全",
                "信任部署 URL 或默认阻断非公网地址",
                "解析校验 私网阻断 禁止凭据和跳转",
                "降低 SSRF 和重定向绕过风险 专网仅允许运维显式开启并配合出口白名单",
            ],
            [
                "恶意文档",
                "依赖 ZIP 库报错或解析前做资源限制",
                "条目数 解压大小 XML 大小 加密和压缩比门禁",
                "在 XML 解析前拒绝 DOCX 解压炸弹 防止内存和 CPU 被恶意文件耗尽",
            ],
        ],
        widths=[1.05, 1.65, 1.85, 2.45],
        font_size=8.1,
    )

    add_heading(document, "十七之三 实现内容", 2)
    add_table(
        document,
        ["区域", "实现内容", "关键边界"],
        [
            [
                "性能场景",
                "二十并发用户覆盖合同创建读取 审批待办和审计查询",
                "每次运行三十秒 首版用于回归基线 不等于生产容量承诺",
            ],
            [
                "CI 证据",
                "生成 Locust CSV HTML JSON Markdown 和 API 日志并上传产物",
                "报告保留提交 SHA 运行 ID 并在阈值失败时阻断合并",
            ],
            [
                "资源成本",
                "统计核心关系存储并计算单合同近似数据库字节",
                "模型默认关闭 因此该基线模型 Token 成本为零",
            ],
            [
                "故障测试",
                "模拟 Redis 暂停 Worker 崩溃后陈旧消息领取和通知端点恢复",
                "失败时不确认消息 不提前标记发布 恢复后只产生一次成功通知",
            ],
            [
                "安全测试",
                "增加 DOCX 解压炸弹 Webhook 私网和日志敏感字段回归",
                "合同正文 JWT Payload 不进入结构化日志输出",
            ],
            [
                "运维交付",
                "新增发布门禁 故障分级 处置 回滚 压测复现和十五分钟演示脚本",
                "所有操作保留 request trace 事件标识 禁止直接清理 Outbox 审计和幂等记录",
            ],
        ],
        widths=[1.2, 3.2, 2.6],
        font_size=8.2,
    )

    add_heading(document, "十七之四 问题难点和处理", 2)
    add_heading(document, "没有 Docker 的本机不能提供可信服务 QPS", 3)
    add_paragraph(
        document,
        "本机只能完成纯 Python 单元和类型门禁 直接记录 TestClient 或领域函数吞吐会绕开网络 数据库和 RLS "
        "因此将真实压测放到带 PostgreSQL 和 Redis 服务的 GitHub CI 首次成功产物再回填 QPS 和分位延迟",
    )
    add_heading(document, "性能测试必须同时保护正确性", 3)
    add_paragraph(
        document,
        "只追求吞吐会掩盖审批冲突和重复通知 本阶段在压测前运行 PostgreSQL 冲突 Outbox 去重 Scheduler "
        "和故障注入测试 报告同时列出冲突拦截 恢复率 重复通知率和越权阻断率",
    )
    add_heading(document, "Webhook 合法协议不代表目标安全", 3)
    add_paragraph(
        document,
        "仅检查 HTTP HTTPS 仍可能访问环回 私网和云元数据地址 本次在发送前解析主机并要求全部地址是公网 "
        "同时拒绝 URL 凭据和 HTTP 跳转 专网场景必须由部署方显式开启并承担出口控制",
    )
    add_heading(document, "DOCX 是 ZIP 容器会产生资源放大", 3)
    add_paragraph(
        document,
        "文件上传大小限制不能阻止高压缩内容在内存展开 解析器先检查条目数量 总解压大小 document XML 大小 "
        "单项压缩比和加密标记 超限时使用稳定业务错误终止而不进入 XML 解析",
    )
    add_heading(document, "CI 基线不能包装为生产容量", 3)
    add_paragraph(
        document,
        "共享 Runner 的硬件和邻居负载会变化 报告强制记录运行环境并明确用途是回归比较 上线前仍需在目标规格 "
        "预发布环境延长测试时间 观察数据库 CPU 连接 Redis 延迟和业务峰值",
    )

    add_heading(document, "十七之五 验证结果", 2)
    add_table(
        document,
        ["验证项", "结果", "证据或限制"],
        [
            ["Ruff", "通过", "API 源码 测试 脚本和 performance 目录无规则错误"],
            ["mypy strict", "通过", "五十二个源码文件无类型错误"],
            ["本地 pytest", "通过", "六十二项通过 十项 PostgreSQL 或 Redis 测试按环境跳过"],
            ["M7 安全回归", "通过", "五项故障与安全测试全部通过"],
            ["PR 七", "已合并", "M6 七项远程检查全部通过并自动合并"],
            [
                "M7 实际 QPS P95 P99",
                "待 PR CI",
                "本机没有 Docker 不虚构数字 首个 M7 CI 产物生成后回填验收报告",
            ],
            ["M7 远程门禁", "待创建", "阶段分支推送后执行质量 集成 性能和仓库全量检查"],
        ],
        widths=[1.7, 1.0, 4.25],
        font_size=8.5,
    )

    add_heading(document, "十七之六 遗留风险和下一步", 2)
    add_bullets(
        document,
        [
            "CI 三十秒基线用于发现回退 不覆盖长时间连接池耗尽和数据库膨胀 上线前需要五分钟以上稳态与峰值测试",
            "DNS 解析校验降低 SSRF 风险 生产仍应使用出口代理 域名白名单和云元数据网络隔离防止 DNS 重绑定",
            "扫描 PDF OCR 和模型提示评测仍属于后续增强 不影响结构化审批履约主链路",
            "M7 PR 首次成功后下载证据产物 将真实 QPS P50 P95 P99 失败率和单合同数据库成本回填报告",
        ],
    )

    return [element for element in document.element.body if element not in before]


def _move_before(elements: list[object], anchor) -> None:
    for element in elements:
        anchor.addprevious(element)


def main() -> None:
    document = Document(DOCUMENT_PATH)
    update_exists = any(
        paragraph.text == "十七 更新记录 UPD013" for paragraph in document.paragraphs
    )
    if not update_exists:
        anchors = {paragraph.text: paragraph for paragraph in document.paragraphs}
        decision_anchor = anchors["十七 当前决策记录"]
        new_elements = _add_update(document)
        _move_before(new_elements, decision_anchor._p)

        decision_anchor.text = "十八 当前决策记录"
        anchors["十八 风险登记"].text = "十九 风险登记"
        anchors["十九 后续更新记录模板"].text = "二十 后续更新记录模板"
        for paragraph in (
            decision_anchor,
            anchors["十八 风险登记"],
            anchors["十九 后续更新记录模板"],
        ):
            for run in paragraph.runs:
                set_run_font(run, size=16, bold=True)

        decisions = _table_with_headers(document, ("编号", "决策", "原因", "重新评估条件"))
        _append_table_row(
            decisions,
            [
                "ADR 037",
                "真实依赖 CI 生成性能证据",
                "本机没有 Docker 不用内存替身吞吐冒充服务 QPS 报告保留提交和运行环境",
                "拥有固定规格预发布集群或独立压测机时重新建立容量基线",
            ],
        )
        _append_table_row(
            decisions,
            [
                "ADR 038",
                "出站通知与压缩文档默认失败关闭",
                "Webhook 目标和 DOCX 容器都属于不可信输入 解析或网络边界异常时先拒绝",
                "引入统一出口代理 沙箱解析服务或可信专网通知平台时重新评估",
            ],
        )

        risks = _table_with_headers(document, ("风险", "影响", "当前控制", "后续验证"))
        _append_table_row(
            risks,
            [
                "共享 CI 性能波动",
                "单次 QPS 或延迟不能代表生产容量",
                "固定剖面 阈值 原始产物 提交 SHA 和运行 ID",
                "预发布环境五分钟以上稳态 峰值和多轮对比",
            ],
        )
        _append_table_row(
            risks,
            [
                "DNS 重绑定或出口绕过",
                "Webhook 可能访问内部服务",
                "发送前解析全部地址 禁止私网 凭据和重定向",
                "生产出口代理 域名白名单和云元数据隔离",
            ],
        )

    _ensure_overview(document)
    _set_metadata(document)
    document.save(DOCUMENT_PATH)
    print(DOCUMENT_PATH.resolve())


if __name__ == "__main__":
    main()

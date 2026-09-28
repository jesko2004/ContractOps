from __future__ import annotations

from pathlib import Path

from build_contractops_update_log import (
    PALE_BLUE,
    add_bullets,
    add_heading,
    add_paragraph,
    add_table,
    prevent_row_split,
    set_cell_margins,
    set_cell_shading,
    set_run_font,
)
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOCUMENT_PATH = Path("architecture/ContractOps-项目更新记录.docx")


def _set_cell(cell, value: str, *, shaded: bool = False, centered: bool = False) -> None:
    cell.text = ""
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if shaded:
        set_cell_shading(cell, PALE_BLUE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(value)
    set_run_font(run, size=8.7)


def _append_table_row(table, values: list[str]) -> None:
    row = table.add_row()
    prevent_row_split(row)
    shaded = (len(table.rows) - 2) % 2 == 1
    for index, value in enumerate(values):
        _set_cell(row.cells[index], value, shaded=shaded, centered=index == 0)


def _table_with_headers(document: Document, headers: tuple[str, ...]):
    for table in document.tables:
        if tuple(cell.text.strip() for cell in table.rows[0].cells) == headers:
            return table
    raise RuntimeError(f"table not found: {headers}")


def _set_cover_metadata(document: Document) -> None:
    header_text = "ContractOps 项目更新记录  版本 1.2"
    for section in document.sections:
        paragraph = section.header.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(header_text)
        set_run_font(run, size=8.5)

    cover = _table_with_headers(document, ("项目字段", "当前内容"))
    values = {"文档版本": "1.2", "最后更新": "2026 年 9 月 28 日"}
    for row in cover.rows[1:]:
        key = row.cells[0].text.strip()
        if key in values:
            _set_cell(row.cells[1], values[key])

    document.paragraphs[4].text = (
        "当前结论  项目主线是合同审批与履约风控后端服务 M0 至 M6 已形成从多租户台账 "
        "审批 履约 可靠事件 审计到文档上传 证据解析和版本差异的后端闭环"
    )
    for run in document.paragraphs[4].runs:
        set_run_font(run)


def _ensure_overview_row(document: Document) -> None:
    overview = _table_with_headers(document, ("编号", "日期", "主题", "结果", "验证状态"))
    if any(row.cells[0].text.strip() == "UPD 012" for row in overview.rows[1:]):
        return
    _append_table_row(
        overview,
        [
            "UPD 012",
            "2026 09 28",
            "实现 M6 文档辅助与版本 Diff",
            "新增预签名上传 完整性校验 PDF DOCX 解析 草稿提示和证据化版本差异",
            "本地 57 项测试 Ruff mypy 通过 数据库迁移待 PR CI",
        ],
    )


def _normalize_m6_decision_number(document: Document) -> None:
    decisions = _table_with_headers(document, ("编号", "决策", "原因", "重新评估条件"))
    for row in decisions.rows[1:]:
        if row.cells[1].text.strip() == "预签名直传加验证后异步解析":
            _set_cell(row.cells[0], "ADR 036", centered=True)
            return


def _normalize_m6_validation_count(document: Document) -> None:
    replacements = {"本地 56 项测试": "本地 57 项测试", "五十六项通过": "五十七项通过"}
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            for old, new in replacements.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)


def _add_update(document: Document) -> list[object]:
    before = set(document.element.body)

    add_heading(document, "十六 更新记录 UPD012", 1)
    add_heading(document, "十六之一 更新目标", 2)
    add_paragraph(
        document,
        "执行实施计划 M6 文档辅助和版本 Diff 建立不经过 API 进程传输大文件的上传路径 "
        "实现 PDF 和 DOCX 解析 金额日期提示 证据查询与版本差异 并保证模型输出只形成草稿提示 "
        "不能直接改变合同 审批 风险或履约状态",
    )

    add_heading(document, "十六之二 候选方案和选择过程", 2)
    add_table(
        document,
        ["决策项", "候选方法", "最终选择", "选择依据"],
        [
            [
                "文件上传",
                "API multipart 中转或对象存储预签名直传",
                "MinIO 预签名 PUT 加完成确认",
                "避免大文件占用 API 内存和连接 上传完成后再验证对象元数据",
            ],
            [
                "解析执行",
                "请求内同步解析或独立异步任务",
                "数据库租约任务加 ingestion worker",
                "解析可重试 可恢复 且不会拖慢业务请求",
            ],
            [
                "DOCX 读取",
                "引入完整办公套件或读取 OOXML",
                "标准库 ZIP XML 读取正文和标题层级",
                "减少运行依赖并保留标题路径 PDF 使用 pypdf 保留页码",
            ],
            [
                "提示落库",
                "直接创建风险义务或保存候选证据",
                "独立 document findings 草稿表",
                "法律和履约决定必须由业务人员确认 模型没有写业务状态的通道",
            ],
            [
                "版本差异",
                "整篇字符串 Diff 或 Chunk 序列 Diff",
                "按内容哈希进行序列匹配",
                "新增 删除 替换都同时返回旧新版本 Chunk 页码和原文",
            ],
        ],
        widths=[1.05, 1.7, 1.75, 2.5],
        font_size=8.2,
    )

    add_heading(document, "十六之三 实现内容", 2)
    add_table(
        document,
        ["区域", "实现内容", "关键边界"],
        [
            [
                "上传协议",
                "新增上传票据和完成确认接口 URL 默认十五分钟失效",
                "只接受 PDF DOCX 和一百 MB 以内文件 对象键包含租户和合同",
            ],
            [
                "完整性校验",
                "完成确认校验对象大小 Worker 再校验 SHA 256",
                "声明与实际内容不符时终止任务 不发布 Chunk",
            ],
            [
                "解析任务",
                "使用 SKIP LOCKED 租约 重试退避和终态失败",
                "Chunk 与提示在同一数据库事务发布 READY 只在成功后设置",
            ],
            [
                "证据模型",
                "Chunk 保存版本 序号 页码 标题路径 原文和哈希",
                "旧版本 Chunk 不覆盖 外键阻止删除仍被引用的证据",
            ],
            [
                "规则提示",
                "从原文提取金额和日期并保存标准化值",
                "每条提示保留 source chunk 页码 原文和置信度",
            ],
            [
                "模型建议",
                "可选调用 OpenAI 兼容模型网关 默认关闭",
                "非法字段 非原文引用或置信度越界会整批丢弃 模型故障降级为空",
            ],
            [
                "版本 Diff",
                "返回 ADDED REMOVED REPLACED 及双侧证据",
                "先校验两个版本属于同一合同并复用租户和部门权限",
            ],
            [
                "运行部署",
                "新增 MinIO 初始化任务 ingestion worker 配置和运行说明",
                "Worker 使用受限 BYPASSRLS 角色 只授权解析所需表",
            ],
        ],
        widths=[1.15, 3.3, 2.55],
        font_size=8.2,
    )

    add_heading(document, "十六之四 问题难点和处理", 2)
    add_heading(document, "完成上传不能信任客户端声明", 3)
    add_paragraph(
        document,
        "客户端提供的文件大小和摘要只能用于建立预期值 完成回调先从对象存储读取实际大小 "
        "Worker 下载后再计算 SHA 256 只有两次校验都通过才进入解析和证据发布",
    )
    add_heading(document, "解析重试不能产生半成品证据", 3)
    add_paragraph(
        document,
        "任务先通过数据库租约独占处理 解析在内存完成后才在一个事务内写入 Chunk 提示和 READY 状态 "
        "失败只更新重试或终态信息 不让查询接口看到部分结果",
    )
    add_heading(document, "DOCX 没有可靠的物理页码", 3)
    add_paragraph(
        document,
        "PDF 保存真实页码 DOCX 受字体和排版环境影响不伪造页码 而是保存标题路径和顺序 "
        "接口明确返回空页码 避免把不可靠定位包装为准确证据",
    )
    add_heading(document, "模型结构化输出仍然不可信", 3)
    add_paragraph(
        document,
        "模型响应必须是受限 JSON 数组 风险或义务引用必须逐字存在于当前 Chunk "
        "任何一项字段越界就整批拒绝 网络失败或解析失败降级为空 且结果只写独立草稿提示表",
    )
    add_heading(document, "版本差异必须保留双侧证据", 3)
    add_paragraph(
        document,
        "只返回合并后的文本会丢失旧合同引用 本次 Diff 以不可变 Chunk 哈希匹配序列 "
        "每个变化同时携带旧侧和新侧的版本内 Chunk 标识 页码 原文和哈希",
    )

    add_heading(document, "十六之五 验证结果", 2)
    add_table(
        document,
        ["验证项", "结果", "证据或限制"],
        [
            ["Ruff", "通过", "API 源码 测试和脚本无规则错误"],
            ["mypy strict", "通过", "五十二个源码文件无类型错误"],
            ["本地 pytest", "通过", "五十七项通过 十项外部数据库或 Redis 测试按环境跳过"],
            ["Alembic 版本图", "通过", "20260928 0007 为唯一 head"],
            ["完整检查脚本", "通过", "Ruff mypy 和 pytest 组合门禁通过"],
            [
                "数据库迁移往返",
                "待 PR CI",
                "本机没有 Docker 可执行程序 由 GitHub 空库升级回滚门禁验证",
            ],
            ["远程 PR 门禁", "待创建", "阶段分支推送后自动执行七项检查并自动合并"],
        ],
        widths=[1.65, 1.0, 4.3],
        font_size=8.6,
    )

    add_heading(document, "十六之六 遗留风险和下一步", 2)
    add_bullets(
        document,
        [
            "扫描版 PDF 当前没有 OCR 文本时不会生成 Chunk "
            "后续根据真实样本决定是否接入 MinerU 或 OCR 服务",
            "DOCX 页码保持为空 后续如业务必须定位物理页需要引入受控渲染服务并建立版式一致性测试",
            "模型辅助默认关闭 上线前需要建立引用准确率 风险召回率和恶意 Prompt 样本评测",
            "下一阶段 M7 完成业务演示 压测 安全回归 运维手册和简历可复现指标",
        ],
    )

    return [element for element in document.element.body if element not in before]


def _move_before(elements: list[object], anchor) -> None:
    for element in elements:
        anchor.addprevious(element)


def main() -> None:
    document = Document(DOCUMENT_PATH)
    update_exists = any(
        paragraph.text == "十六 更新记录 UPD012" for paragraph in document.paragraphs
    )
    if not update_exists:
        anchors = {paragraph.text: paragraph for paragraph in document.paragraphs}
        decision_anchor = anchors["十六 当前决策记录"]
        new_elements = _add_update(document)
        _move_before(new_elements, decision_anchor._p)

        decision_anchor.text = "十七 当前决策记录"
        anchors["十七 风险登记"].text = "十八 风险登记"
        anchors["十八 后续更新记录模板"].text = "十九 后续更新记录模板"
        for paragraph in (
            decision_anchor,
            anchors["十七 风险登记"],
            anchors["十八 后续更新记录模板"],
        ):
            for run in paragraph.runs:
                set_run_font(run, size=16, bold=True)

        decisions = _table_with_headers(document, ("编号", "决策", "原因", "重新评估条件"))
        _append_table_row(
            decisions,
            [
                "ADR 036",
                "预签名直传加验证后异步解析",
                "隔离大文件传输和耗时解析 并在发布证据前验证大小与摘要",
                "对象存储协议或安全边界发生变化时重新评估",
            ],
        )

        risks = _table_with_headers(document, ("风险", "影响", "当前控制", "后续验证"))
        _append_table_row(
            risks,
            [
                "上传声明与对象内容不一致",
                "错误文件进入解析或证据链",
                "完成回调校验大小 Worker 校验 SHA 256",
                "对象替换 分段上传和摘要冲突测试",
            ],
        )
        _append_table_row(
            risks,
            [
                "模型建议引用伪造",
                "无依据风险或义务候选误导人工",
                "默认关闭 严格字段白名单 原文子串校验 整批失败关闭",
                "恶意 Prompt 和引用准确率评测集",
            ],
        )

    _ensure_overview_row(document)
    _normalize_m6_decision_number(document)
    _normalize_m6_validation_count(document)
    _set_cover_metadata(document)
    document.save(DOCUMENT_PATH)
    print(DOCUMENT_PATH.resolve())


if __name__ == "__main__":
    main()
